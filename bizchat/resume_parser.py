"""Local resume extraction. Outbound text is reconstructed from allowlisted facts.

Reconstruction deliberately drops personal headers, employers, institutions and
free-form prose: regex-only redaction cannot reliably identify arbitrary names.
"""

from __future__ import annotations

import io
import re
import zipfile

from bizchat.job_classifier import classify_job
from bizchat.skills import extract_skills


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 1000
MAX_TEXT_CHARS = 200_000

_YEAR_VALUE = r"(\d{1,2}(?:\.\d)?)"
_SECTION_HEADINGS = r"education|academic\s+background|qualifications?|experience|employment|projects?|certifications?|skills"


class _SafeExtractionError(ValueError):
    pass


def _extract_experience(text: str) -> float | None:
    """Extract only explicitly labelled or resume-style experience durations."""
    patterns = (
        # 3 years of professional experience / 3+ yrs experience
        rf"\b{_YEAR_VALUE}\s*\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:professional\s+|work\s+|industry\s+|relevant\s+)?experience\b",
        # experience: 4 years / professional experience 6 years
        rf"\b(?:professional\s+|total\s+|relevant\s+|work\s+|industry\s+)?experience\s*:?\s*{_YEAR_VALUE}\s*\+?\s*(?:years?|yrs?)\b",
    )
    values = [float(value) for pattern in patterns for value in re.findall(pattern, text, re.I)]
    # A standalone duration is common in compact resume headers, but only accept
    # it when the whole line is the duration; prose such as project timelines is
    # intentionally excluded.
    values.extend(float(value) for value in re.findall(
        rf"^\s*(?:total\s+)?{_YEAR_VALUE}\s*\+?\s*(?:years?|yrs?)\s*$",
        text, re.I | re.M,
    ))
    valid = [value for value in values if value <= 60]
    return max(valid, default=None)


def _education_lines(text: str) -> list[str]:
    section = re.search(
        rf"(?ims)^\s*(?:education|academic\s+background|qualifications?)\s*:?\s*$\n?(.*?)(?=^\s*(?:{_SECTION_HEADINGS})\s*:?\s*$|\Z)",
        text,
    )
    lines = section.group(1).splitlines() if section else []
    # Resumes without an Education heading still commonly put a degree on its
    # own line. Keep those candidates line-scoped to avoid project prose.
    lines.extend(text.splitlines())
    return lines


def _extract_education(text: str) -> list[str]:
    degrees = (
        ("Doctorate", r"\b(?:ph\.?\s*d\.?|doctorate|doctor\s+of\s+philosophy)\b"),
        ("Master's degree", r"\b(?:m\.?\s*(?:tech|sc|s)\.?|master(?:'s|s)?(?:\s+degree)?|master\s+of)\b"),
        ("Bachelor's degree", r"\b(?:b\.?\s*(?:tech|sc|e)\.?|bachelor(?:'s|s)?(?:\s+degree)?|bachelor\s+of)\b"),
    )
    found = []
    section_match = re.search(
        rf"(?ims)^\s*(?:education|academic\s+background|qualifications?)\s*:?\s*$\n?(.*?)(?=^\s*(?:{_SECTION_HEADINGS})\s*:?\s*$|\Z)",
        text,
    )
    section_lines = set(section_match.group(1).splitlines()) if section_match else set()
    for line in _education_lines(text):
        if not line.strip():
            continue
        for label, pattern in degrees:
            if not re.search(pattern, line, re.I):
                continue
            # Degree-led lines and explicit degree fields are strong signals;
            # otherwise require institution/graduation context on the same line.
            contextual = re.search(
                r"^\s*(?:degree\s*[:\-]|(?:ph\.?\s*d\.?|doctorate|m\.?\s*(?:tech|sc|s)\.?|master|b\.?\s*(?:tech|sc|e)\.?|bachelor))\b"
                r"|\b(?:university|college|school|institute|graduat(?:ed|ion)|class\s+of|major)\b"
                r"|\b(?:19|20)\d{2}\b",
                line,
                re.I,
            )
            # Outside an Education section, require a degree-led/degree-field
            # line; an arbitrary project sentence mentioning a degree is not
            # sufficient even if it mentions an institution elsewhere.
            if line not in section_lines:
                contextual = contextual and re.search(
                    r"^\s*(?:degree\s*[:\-]|(?:ph\.?\s*d\.?|doctorate|m\.?\s*(?:tech|sc|s)\.?|master|b\.?\s*(?:tech|sc|e)\.?|bachelor))\b"
                    r".*(?:\b(?:degree|in|of|university|college|school|institute|graduat(?:ed|ion))\b|\b(?:19|20)\d{2}\b)",
                    line,
                    re.I,
                )
            if contextual and label not in found:
                found.append(label)
    return found


def parse_resume(text: str) -> dict:
    if not isinstance(text, str):
        raise ValueError("Resume text must be a string.")
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError("Resume text is too long.")
    skills = extract_skills(text)
    titles = []
    for line in text.splitlines():
        role = classify_job({"title": line})["role_category"]
        if role != "Unknown" and role not in titles:
            titles.append(role)
    experience = _extract_experience(text)
    education = _extract_education(text)
    projects = []
    project_section = re.search(r"(?ims)^\s*(?:personal\s+|selected\s+)?projects?\s*:?\s*\n(.*?)(?=^\s*(?:education|experience|employment|certifications?|skills)\s*:?\s*$|\Z)", text)
    if project_section:
        project_skills = extract_skills(project_section.group(1))
        if project_skills:
            projects = ["Project technologies: " + ", ".join(project_skills)]
    safe = []
    if skills:
        safe.append("Skills: " + ", ".join(skills))
    if titles:
        safe.append("Roles: " + ", ".join(titles))
    if experience is not None:
        safe.append(f"Explicit experience: {experience:g} years")
    if education:
        safe.append("Education: " + ", ".join(education))
    safe.extend(projects)
    return {"skills": skills, "titles": titles, "years_experience": experience,
            "education": education, "projects": projects, "redacted_text": "\n".join(safe)}


def extract_resume_file(data: bytes, filename: str) -> str:
    """Extract PDF/DOCX, raising safe ValueErrors without parser internals."""
    if not isinstance(data, bytes) or not data:
        raise ValueError("Resume file is empty or invalid.")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("Resume file exceeds the 5 MiB limit.")
    extension = str(filename).rsplit(".", 1)[-1].lower()
    if extension not in {"pdf", "docx"}:
        raise ValueError("Unsupported resume file. Use PDF or DOCX.")
    try:
        if extension == "pdf":
            if not data.startswith(b"%PDF-"):
                raise _SafeExtractionError("Invalid PDF file.")
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise _SafeExtractionError("Encrypted PDF files are not supported.")
            if len(reader.pages) > 100:
                raise _SafeExtractionError("Resume PDF has too many pages.")
            chunks = []
            size = 0
            for page in reader.pages:
                chunk = page.extract_text() or ""
                size += len(chunk)
                if size > MAX_TEXT_CHARS:
                    raise _SafeExtractionError("Extracted resume text is too long.")
                chunks.append(chunk)
            text = "\n".join(chunks)
        else:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if (len(entries) > MAX_ARCHIVE_ENTRIES
                        or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES
                        or any(entry.flag_bits & 1 or entry.file_size / max(entry.compress_size, 1) > 200 for entry in entries)):
                    raise _SafeExtractionError("DOCX archive exceeds safe expansion limits or is encrypted.")
                if "word/document.xml" not in archive.namelist() or "[Content_Types].xml" not in archive.namelist():
                    raise _SafeExtractionError("Invalid DOCX file.")
            from docx import Document

            document = Document(io.BytesIO(data))
            text = "\n".join([p.text for p in document.paragraphs] +
                             [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    except ImportError:
        raise ValueError("Resume file extraction dependency is unavailable.") from None
    except _SafeExtractionError:
        raise
    except Exception:
        raise ValueError("Resume file could not be read. Use a valid, unencrypted PDF or DOCX.") from None
    if not text.strip():
        raise ValueError("No readable text found. Scanned resumes require OCR before upload.")
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError("Extracted resume text is too long.")
    return text.strip()
